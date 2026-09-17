"""
export_excel.py
Penyusunan dataset ke DataFrame dan penyimpanan ke file Excel (.xlsx) beserta formatting styling.
"""

import os
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def save_to_excel(records: list[dict], output_path: str, merge_existing: bool = False):
    """
    Menyimpan hasil scraping dan NLP ke file Excel dengan styling:
    - Kolom wajib: Tanggal, Title, Link Website, Media Name, Tone, Spokesperson 1, Spokesperson 2, Unit Eselon, Terkait Kemenperin, Keywords, Sumber Data
    - Mendukung merge_existing=True untuk menggabungkan dengan file yang sudah ada (dedup URL dan judul)
    - Header bold dan lebar kolom auto-fit
    - Kolom Link Website berupa hyperlink aktif
    - Kolom Tanggal diformat sebagai tanggal Excel (datetime tanpa timezone)
    """
    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    final_input_records = list(records)
    if merge_existing and os.path.exists(output_path):
        try:
            existing_df = pd.read_excel(output_path)
            if not existing_df.empty:
                from fetch_news import dedup_by_link, dedup_by_title
                combined_raw = existing_df.to_dict(orient="records") + final_input_records
                link_deduped = dedup_by_link(combined_raw)
                final_input_records = dedup_by_title(link_deduped, threshold=85)
        except Exception as e:
            print(f"  [WARNING] Gagal merge dengan file yang ada ({output_path}): {e}")

    columns = [
        "Tanggal",
        "Title",
        "Link Website",
        "Media Name",
        "Tone",
        "Spokesperson 1",
        "Spokesperson 2",
        "Unit Eselon",
        "Terkait Kemenperin",
        "Keywords",
        "Sumber Data",
    ]
    df = pd.DataFrame(final_input_records)
    if not df.empty:
        # 1. Normalisasi nama kolom dari berbagai variasi key dictionary pipeline
        mappings = {
            "Title": ["Title", "title", "judul"],
            "Link Website": ["Link Website", "link", "url", "resolved_url", "link_website"],
            "Tanggal": ["Tanggal", "published", "tanggal", "pub_date"],
            "Media Name": ["Media Name", "media_name", "source", "media"],
            "Tone": ["Tone", "tone", "sentimen"],
            "Spokesperson 1": ["Spokesperson 1", "spokesperson_1", "sp1"],
            "Spokesperson 2": ["Spokesperson 2", "spokesperson_2", "sp2"],
            "Unit Eselon": ["Unit Eselon", "unit_eselon", "unit"],
            "Terkait Kemenperin": ["Terkait Kemenperin", "terkait_kemenperin"],
            "Keywords": ["Keywords", "keywords", "Keyword", "keyword"],
            "Sumber Data": ["Sumber Data", "sumber_data", "source_data"],
        }

        for target_col, alt_keys in mappings.items():
            if target_col not in df.columns:
                df[target_col] = None
            for alt in alt_keys:
                if alt in df.columns and alt != target_col:
                    df[target_col] = df[target_col].fillna(df[alt])

        # Fallback default untuk Sumber Data
        if "Sumber Data" not in df.columns or df["Sumber Data"].isna().all():
            df["Sumber Data"] = "RSS"
        df["Sumber Data"] = df["Sumber Data"].replace("", "RSS").fillna("RSS")

        # Pastikan kolom bertipe string/object agar aman saat assignment
        text_cols = [c for c in columns if c != "Tanggal"]
        for col in text_cols:
            if col not in df.columns:
                df[col] = None
            df[col] = df[col].astype("object")

        # 2. Auto-enrichment NLP & Deteksi Ketat Kemenperin
        from entity_mapper import find_spokespersons
        from relevance_filter import is_kemenperin_related
        from sentiment import classify_tone

        for idx in range(len(df)):
            t = str(df.at[idx, "Title"] or "")
            txt = str(df.at[idx, "text"] or "") if "text" in df.columns else ""
            kw = str(df.at[idx, "Keywords"] or "")
            sd = str(df.at[idx, "Sumber Data"] or "")

            combined = f"{t} {txt}"

            # 1. Spokesperson 1 & 2 serta Unit Eselon (diurutkan berdasar kemunculan pertama)
            sp1, sp2, unit = find_spokespersons(combined)
            df.at[idx, "Spokesperson 1"] = sp1
            df.at[idx, "Spokesperson 2"] = sp2
            df.at[idx, "Unit Eselon"] = unit if unit else "-"

            # 2. Terkait Kemenperin (KETAT & SPESIFIK)
            is_rel = is_kemenperin_related(title=t, text=txt, keyword=kw, sumber_data=sd)
            df.at[idx, "Terkait Kemenperin"] = "Ya" if is_rel else "Tidak"

            # 3. Tone
            if pd.isna(df.at[idx, "Tone"]) or str(df.at[idx, "Tone"]).strip() in ["", "nan", "None"]:
                df.at[idx, "Tone"] = classify_tone(txt, t)

        # Pembersihan nilai kosong agar tidak tersimpan sebagai string literal 'nan'
        if "Tone" in df.columns:
            df["Tone"] = df["Tone"].replace(["nan", "None", None, ""], "Netral").fillna("Netral")
        if "Unit Eselon" in df.columns:
            df["Unit Eselon"] = df["Unit Eselon"].replace(["nan", "None", None, ""], "-").fillna("-")
        if "Spokesperson 1" in df.columns:
            df["Spokesperson 1"] = df["Spokesperson 1"].replace(["nan", "None", None], "").fillna("")
        if "Spokesperson 2" in df.columns:
            df["Spokesperson 2"] = df["Spokesperson 2"].replace(["nan", "None", None], "").fillna("")
        if "Terkait Kemenperin" in df.columns:
            df["Terkait Kemenperin"] = df["Terkait Kemenperin"].replace(["nan", "None", None, ""], "Tidak").fillna("Tidak")

        # Pastikan seluruh kolom wajib ada dan urut
        for col in columns:
            if col not in df.columns:
                df[col] = ""
        df = df[columns]
    else:
        df = pd.DataFrame(columns=columns)



    # Konversi field Tanggal ke datetime Excel tanpa timezone
    if "Tanggal" in df.columns and not df.empty:
        import dateparser

        def _parse_date(val):
            if pd.isna(val) or not val:
                return None
            if isinstance(val, pd.Timestamp):
                return val.tz_localize(None) if val.tzinfo else val
            try:
                dt = pd.to_datetime(val, utc=True)
                if not pd.isna(dt):
                    return dt.tz_localize(None) if dt.tzinfo else dt
            except Exception:
                pass
            try:
                parsed = dateparser.parse(str(val))
                if parsed:
                    return pd.Timestamp(parsed)
            except Exception:
                pass
            return None

        df["Tanggal"] = df["Tanggal"].apply(_parse_date)

    def _write_excel(target_writer):
        df.to_excel(target_writer, index=False, sheet_name="Berita")
        ws = target_writer.sheets["Berita"]

        # Format header bold dan lebar kolom auto-fit
        for col_idx, col_name in enumerate(columns, start=1):
            ws.cell(row=1, column=col_idx).font = Font(bold=True)
            col_lengths = [len(str(v)) for v in df[col_name].dropna() if str(v).strip()]
            max_len = max(max(col_lengths) if col_lengths else 0, len(col_name))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

        # Format kolom Tanggal sebagai Date Excel
        date_col_idx = columns.index("Tanggal") + 1
        for row_idx in range(2, len(df) + 2):
            cell = ws.cell(row=row_idx, column=date_col_idx)
            if cell.value is not None:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"

        # Buat Link Website jadi hyperlink aktif (href) dengan warna biru dan garis bawah
        link_col_idx = columns.index("Link Website") + 1
        link_font = Font(color="0563C1", underline="single")
        for row_idx in range(2, len(df) + 2):
            cell = ws.cell(row=row_idx, column=link_col_idx)
            val = str(cell.value or "").strip()
            if val and (val.startswith("http://") or val.startswith("https://")):
                cell.hyperlink = val
                cell.font = link_font

    try:
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            _write_excel(writer)
        final_path = output_path
    except PermissionError:
        base, ext = os.path.splitext(output_path)
        final_path = f"{base}_terbaru{ext}"
        print(f"[WARNING] File '{output_path}' sedang dibuka/dikunci oleh Microsoft Excel.")
        print(f"          Menyimpan ke file alternatif: '{final_path}'")
        with pd.ExcelWriter(final_path, engine="openpyxl") as writer:
            _write_excel(writer)

    return final_path
